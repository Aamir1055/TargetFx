"""File exports (CSV / XLSX / PDF) for tables the assistant hands out.

Files live next to the charts in files/ under an unguessable id for an
hour and are served by GET /ai/files/{id}.{ext}; the web page offers a
download link, the Telegram bot sends them as documents.
"""
from __future__ import annotations

import csv
import io
import re
import secrets
from datetime import datetime
from typing import Any

from .charts import FILES_DIR, _cleanup

FORMATS = ("csv", "xlsx", "pdf")
MEDIA_TYPES = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "png": "image/png",
}


def safe_filename(name: str, ext: str) -> str:
    base = re.sub(r"[^A-Za-z0-9._ -]+", "_", (name or "export").strip()).strip("._ ") or "export"
    base = re.sub(r"\.(csv|xlsx|pdf)$", "", base, flags=re.I)
    return f"{base[:80]}.{ext}"


def _cell(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, float):
        return round(v, 2)
    if isinstance(v, (dict, list)):
        return str(v)
    return v


def _rows(columns: list[str], rows: list[Any]) -> list[list[Any]]:
    """Accept rows as lists (in column order) or dicts (keyed by column)."""
    out = []
    for r in rows:
        if isinstance(r, dict):
            out.append([_cell(r.get(c)) for c in columns])
        else:
            out.append([_cell(x) for x in list(r)[: len(columns)]] + [""] * max(0, len(columns) - len(r)))
    return out


def _build_csv(title: str, columns: list[str], rows: list[list[Any]], totals: dict[str, Any] | None) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(columns)
    w.writerows(rows)
    if totals:
        w.writerow([])
        w.writerow(["Total"] + [_cell(totals.get(c, "")) for c in columns[1:]])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 correctly


def _build_xlsx(title: str, columns: list[str], rows: list[list[Any]], totals: dict[str, Any] | None,
                subtitle: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = (re.sub(r"[\[\]:*?/\\]", " ", title)[:28] or "Report").strip() or "Report"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = subtitle or f"Generated {datetime.now():%d %b %Y %H:%M}"
    ws["A2"].font = Font(color="6B7280", size=10)
    head_row = 4
    fill = PatternFill("solid", fgColor="1F2937")
    thin = Side(style="thin", color="E5E7EB")
    for i, c in enumerate(columns, start=1):
        cell = ws.cell(row=head_row, column=i, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r_i, r in enumerate(rows, start=head_row + 1):
        for c_i, v in enumerate(r, start=1):
            cell = ws.cell(row=r_i, column=c_i, value=v)
            cell.border = Border(bottom=thin)
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                cell.number_format = "#,##0.00" if isinstance(v, float) else "#,##0"
                cell.alignment = Alignment(horizontal="right")
    if totals:
        r_i = head_row + len(rows) + 1
        ws.cell(row=r_i, column=1, value="Total").font = Font(bold=True)
        for c_i, c in enumerate(columns[1:], start=2):
            v = _cell(totals.get(c, ""))
            cell = ws.cell(row=r_i, column=c_i, value=v)
            cell.font = Font(bold=True)
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
    for i, c in enumerate(columns, start=1):
        width = max([len(str(c))] + [len(str(r[i - 1])) for r in rows[:500]] + [8])
        ws.column_dimensions[get_column_letter(i)].width = min(48, width + 2)
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(columns))}{head_row + max(1, len(rows))}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _build_pdf(title: str, columns: list[str], rows: list[list[Any]], totals: dict[str, Any] | None,
               subtitle: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    page = landscape(A4) if len(columns) > 6 else A4
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=page, leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title=title)
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]),
             Paragraph(subtitle or f"Generated {datetime.now():%d %b %Y %H:%M}",
                       ParagraphStyle("sub", parent=styles["Normal"], textColor=colors.HexColor("#6B7280"), fontSize=9)),
             Spacer(1, 6 * mm)]

    def fmt(v: Any) -> str:
        if isinstance(v, float):
            return f"{v:,.2f}"
        if isinstance(v, int) and not isinstance(v, bool):
            return f"{v:,}"
        return str(v)

    data = [columns] + [[fmt(v) for v in r] for r in rows]
    if totals:
        data.append(["Total"] + [fmt(_cell(totals.get(c, ""))) for c in columns[1:]])
    numeric_cols = [i for i in range(len(columns))
                    if rows and all(isinstance(r[i], (int, float)) and not isinstance(r[i], bool) for r in rows[:50])]
    t = Table(data, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for i in numeric_cols:
        style.append(("ALIGN", (i, 1), (i, -1), "RIGHT"))
    if totals:
        style += [("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                  ("LINEABOVE", (0, -1), (-1, -1), 0.8, colors.HexColor("#111827"))]
    t.setStyle(TableStyle(style))
    story.append(t)
    doc.build(story)
    return out.getvalue()


def export_table(fmt: str, title: str, columns: list[str], rows: list[Any], *, totals: dict[str, Any] | None = None,
                 subtitle: str = "", filename: str = "") -> dict[str, Any]:
    fmt = (fmt or "xlsx").lower().strip()
    if fmt not in FORMATS:
        raise ValueError(f"format must be one of {', '.join(FORMATS)}")
    if not columns:
        raise ValueError("columns must not be empty")
    columns = [str(c) for c in columns[:60]]
    body = _rows(columns, rows[:50000])
    if fmt == "csv":
        data = _build_csv(title, columns, body, totals)
    elif fmt == "xlsx":
        data = _build_xlsx(title, columns, body, totals, subtitle)
    else:
        data = _build_pdf(title, columns, body, totals, subtitle)

    FILES_DIR.mkdir(exist_ok=True)
    file_id = secrets.token_urlsafe(18)
    path = FILES_DIR / f"{file_id}.{fmt}"
    path.write_bytes(data)
    _cleanup()
    return {"file_id": file_id, "filename": safe_filename(filename or title, fmt), "format": fmt,
            "rows": len(body), "bytes": len(data), "url": f"/ai/files/{file_id}.{fmt}"}
